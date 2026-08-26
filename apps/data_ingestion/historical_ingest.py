import os
import sys
import datetime
from decimal import Decimal
import pandas as pd
import openpyxl
from django.db import transaction
from apps.inventory.models import (
    ProductSKU, Category, ProductionOrderMovement,
    SalesOrderMovement, SKUClassification
)


def ingest_maestro_articulos(file_path):
    """
    Ingesta o actualiza los 20.403 SKUs del Maestro Articulos v7.
    Carga Fecha Creacion, Costo Unitario, Categoria, Clase, Familia, SubFamilia, Modelo.
    """
    print(f"\n[ETL S&OP] Iniciando ingesta de Maestro de Articulos: {file_path}", flush=True)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(max_row=1))]
    col_map = {h: idx for idx, h in enumerate(headers)}

    code_idx     = col_map.get('ItemCode')
    name_idx     = col_map.get('ItemName')
    cost_idx     = col_map.get('$Costo UN')
    clase_idx    = col_map.get('Clase')
    fam_idx      = col_map.get('Familia')
    subfam_idx   = col_map.get('SubFamilia')
    mod_idx      = col_map.get('Modelo')
    grupo_idx    = col_map.get('Nombre Grupo')
    cat_idx      = col_map.get('Categoria')
    fec_crea_idx = col_map.get('Fec.Creacion')
    activo_idx   = col_map.get('Activo')

    # Pre-crear categorías
    categories_cache = {c.name: c for c in Category.objects.all()}

    products_to_create = []
    products_to_update = []
    existing_products = {p.code: p for p in ProductSKU.objects.all()}

    count = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_code = str(row[code_idx]).strip() if code_idx is not None and row[code_idx] is not None else ''
        if not raw_code or raw_code.lower() == 'none' or raw_code == '0':
            continue

        count += 1
        name = str(row[name_idx] or f'SKU {raw_code}').strip()[:250]
        grupo_name = str(row[grupo_idx] or 'General').strip()[:100]
        clase = str(row[clase_idx] or '').strip()[:100]
        familia = str(row[fam_idx] or '').strip()[:100]
        subfamilia = str(row[subfam_idx] or '').strip()[:100]
        modelo = str(row[mod_idx] or '').strip()[:100]
        sap_cat = str(row[cat_idx] or '').strip()[:100]
        is_act = str(row[activo_idx] or 'Y').strip().upper() == 'Y'

        # Costo unitario
        cost_val = 0.0
        try:
            if cost_idx is not None and row[cost_idx] is not None:
                cost_val = max(0.0, float(row[cost_idx]))
        except (ValueError, TypeError):
            cost_val = 0.0

        # Fecha creación SAP
        fec_creacion = None
        if fec_crea_idx is not None and row[fec_crea_idx] is not None:
            raw_date = row[fec_crea_idx]
            if isinstance(raw_date, (datetime.date, datetime.datetime)):
                fec_creacion = raw_date.date() if isinstance(raw_date, datetime.datetime) else raw_date
            elif isinstance(raw_date, str) and len(raw_date) >= 10:
                try:
                    fec_creacion = datetime.datetime.strptime(raw_date[:10], '%Y-%m-%d').date()
                except ValueError:
                    pass

        # Categoría
        cat_obj = categories_cache.get(grupo_name)
        if not cat_obj and grupo_name:
            cat_obj, _ = Category.objects.get_or_create(
                code=f"GRP-{grupo_name[:80]}",
                defaults={'name': grupo_name, 'gerencia': 'Gerencia Repuestos PESCO'}
            )
            categories_cache[grupo_name] = cat_obj

        existing = existing_products.get(raw_code)
        p = ProductSKU(
            code=raw_code,
            description=name,
            category=cat_obj,
            unit_cost_usd=Decimal(str(cost_val)) if cost_val > 0 else (existing.unit_cost_usd if existing else Decimal('0.00')),
            clase=clase,
            familia=familia,
            subfamilia=subfamilia,
            modelo=modelo,
            sap_categoria=sap_cat,
            fec_creacion_sap=fec_creacion,
            is_active=is_act
        )

        if existing:
            products_to_update.append(p)
        else:
            products_to_create.append(p)

    wb.close()

    with transaction.atomic():
        if products_to_create:
            ProductSKU.objects.bulk_create(products_to_create, batch_size=5000, ignore_conflicts=True)
        if products_to_update:
            ProductSKU.objects.bulk_update(
                products_to_update,
                ['description', 'category', 'unit_cost_usd', 'clase', 'familia', 'subfamilia', 'modelo', 'sap_categoria', 'fec_creacion_sap', 'is_active'],
                batch_size=5000
            )

    print(f"[ETL S&OP] Maestro Articulos finalizado: {len(products_to_create)} creados, {len(products_to_update)} actualizados (Total: {count})", flush=True)
    return count


def ingest_detalle_ofs(file_path):
    """
    Ingesta los movimientos de consumo de componentes en Ordenes de Fabricacion (OF).
    Filtra estrictamente por Estado == 'Cerrada'.
    """
    print(f"\n[ETL S&OP] Iniciando ingesta de OFs Cerradas: {file_path}", flush=True)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(max_row=1))]
    col_map = {h: idx for idx, h in enumerate(headers)}

    of_num_idx   = col_map.get('Orden Fabricacion')
    status_idx   = col_map.get('Estado')
    start_idx    = col_map.get('Fec.Inicio')
    end_idx      = col_map.get('Fec.Termino')
    wh_idx       = col_map.get('Almacén', col_map.get('Almacen', col_map.get('Bodega')))
    comp_idx     = col_map.get('Cod.Componente')
    cant_ent_idx = col_map.get('Cant.Entregada')
    cost_idx     = col_map.get('Costo $ Real', col_map.get('Costo $'))

    # Cache de SKUs existentes
    existing_skus = set(ProductSKU.objects.values_list('code', flat=True))

    with transaction.atomic():
        ProductionOrderMovement.objects.all().delete()

    records = []
    batch_size = 10000
    total_saved = 0
    total_processed = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        total_processed += 1
        st = str(row[status_idx] or '').strip()
        if st.lower() != 'cerrada':
            continue

        raw_comp = str(row[comp_idx] or '').strip()
        if not raw_comp or raw_comp not in existing_skus:
            continue

        of_num = str(row[of_num_idx] or '').strip()[:50]
        wh = str(row[wh_idx] or '').strip()[:50]

        # Cantidad entregada
        cant_ent = 0.0
        try:
            if cant_ent_idx is not None and row[cant_ent_idx] is not None:
                cant_ent = float(row[cant_ent_idx])
        except (ValueError, TypeError):
            cant_ent = 0.0

        if cant_ent <= 0:
            continue

        # Costo real
        cost_real = 0.0
        try:
            if cost_idx is not None and row[cost_idx] is not None:
                cost_real = max(0.0, float(row[cost_idx]))
        except (ValueError, TypeError):
            cost_real = 0.0

        # Fechas
        f_start = None
        if start_idx is not None and row[start_idx] is not None:
            raw_s = row[start_idx]
            if isinstance(raw_s, (datetime.date, datetime.datetime)):
                f_start = raw_s.date() if isinstance(raw_s, datetime.datetime) else raw_s

        f_end = None
        if end_idx is not None and row[end_idx] is not None:
            raw_e = row[end_idx]
            if isinstance(raw_e, (datetime.date, datetime.datetime)):
                f_end = raw_e.date() if isinstance(raw_e, datetime.datetime) else raw_e

        records.append(ProductionOrderMovement(
            of_number=of_num,
            product_id=raw_comp,
            quantity_delivered=Decimal(str(cant_ent)),
            warehouse=wh,
            start_date=f_start,
            end_date=f_end,
            status='Cerrada',
            costo_real=Decimal(str(cost_real))
        ))

        if len(records) >= batch_size:
            with transaction.atomic():
                ProductionOrderMovement.objects.bulk_create(records, batch_size=5000)
            total_saved += len(records)
            print(f"[ETL S&OP] OFs Cerradas guardadas: {total_saved:,} registros...", flush=True)
            records = []

    if records:
        with transaction.atomic():
            ProductionOrderMovement.objects.bulk_create(records, batch_size=5000)
        total_saved += len(records)

    wb.close()
    print(f"[ETL S&OP] OFs Cerradas finalizado: {total_saved:,} movimientos guardados (de {total_processed:,} filas leidas)", flush=True)
    return total_saved


def ingest_detalle_pedidos(file_path):
    """
    Ingesta los movimientos de despacho a clientes en Pedidos de Ventas (PC).
    Filtra estrictamente por Estado == 'Cerrado' y Cancelado != 'Y'.
    """
    print(f"\n[ETL S&OP] Iniciando ingesta de Pedidos Clientes Cerrados: {file_path}", flush=True)
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(max_row=1))]
    col_map = {h: idx for idx, h in enumerate(headers)}

    order_idx    = col_map.get('Nro.Pedido')
    status_idx   = col_map.get('Estado')
    canc_idx     = col_map.get('Cancelado')
    f_emi_idx    = col_map.get('Fecha Emision')
    f_ent_idx    = col_map.get('Fecha Entrega')
    prod_idx     = col_map.get('Producto')
    cant_idx     = col_map.get('Cantidad')
    wh_idx       = col_map.get('almacen PC', col_map.get('Bodega'))
    client_idx   = col_map.get('Razon Social', col_map.get('Cod.Cliente'))

    existing_skus = set(ProductSKU.objects.values_list('code', flat=True))

    with transaction.atomic():
        SalesOrderMovement.objects.all().delete()

    records = []
    batch_size = 10000
    total_saved = 0
    total_processed = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        total_processed += 1
        st = str(row[status_idx] or '').strip()
        if st.lower() != 'cerrado':
            continue

        canc = str(row[canc_idx] or '').strip().lower()
        if canc in ['si', 'y', 'yes', 'true', '1']:
            continue

        raw_prod = str(row[prod_idx] or '').strip()
        if not raw_prod or raw_prod not in existing_skus:
            continue

        order_num = str(row[order_idx] or '').strip()[:50]
        wh = str(row[wh_idx] or '').strip()[:50]
        client = str(row[client_idx] or '').strip()[:200]

        # Cantidad despachada
        cant = 0.0
        try:
            if cant_idx is not None and row[cant_idx] is not None:
                cant = float(row[cant_idx])
        except (ValueError, TypeError):
            cant = 0.0

        if cant <= 0:
            continue

        # Fechas
        f_emi = None
        if f_emi_idx is not None and row[f_emi_idx] is not None:
            raw_e = row[f_emi_idx]
            if isinstance(raw_e, (datetime.date, datetime.datetime)):
                f_emi = raw_e.date() if isinstance(raw_e, datetime.datetime) else raw_e

        f_ent = None
        if f_ent_idx is not None and row[f_ent_idx] is not None:
            raw_d = row[f_ent_idx]
            if isinstance(raw_d, (datetime.date, datetime.datetime)):
                f_ent = raw_d.date() if isinstance(raw_d, datetime.datetime) else raw_d

        records.append(SalesOrderMovement(
            order_number=order_num,
            product_id=raw_prod,
            quantity=Decimal(str(cant)),
            warehouse=wh,
            emission_date=f_emi,
            delivery_date=f_ent,
            status='Cerrado',
            client_name=client
        ))

        if len(records) >= batch_size:
            with transaction.atomic():
                SalesOrderMovement.objects.bulk_create(records, batch_size=5000)
            total_saved += len(records)
            print(f"[ETL S&OP] Pedidos Clientes guardados: {total_saved:,} registros...", flush=True)
            records = []

    if records:
        with transaction.atomic():
            SalesOrderMovement.objects.bulk_create(records, batch_size=5000)
        total_saved += len(records)

    wb.close()
    print(f"[ETL S&OP] Pedidos Clientes finalizado: {total_saved:,} movimientos guardados (de {total_processed:,} filas leidas)", flush=True)
    return total_saved

